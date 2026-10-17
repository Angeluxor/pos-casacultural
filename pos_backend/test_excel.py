import openpyxl
wb = openpyxl.load_workbook('data/Sistema_Ventas_MASTER.xlsx', read_only=True, data_only=True)
print('Sheets:', wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    row1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    print(s, ':', row1)
