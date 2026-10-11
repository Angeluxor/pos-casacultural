from datetime import datetime
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.timezone import BOGOTA_TZ

PRIMARIO = colors.HexColor("#892325")
PRIMARIO_FILL = "892325"
SUAVE_FILL = "F5ECDB"
CELDA_SUAVE = colors.HexColor("#f5ecdb")


def _sin_zona(valor: datetime) -> datetime:
    if valor.tzinfo is not None:
        return valor.astimezone(BOGOTA_TZ).replace(tzinfo=None)
    return valor


def _pdf_celda(valor, es_money: bool) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, float):
        return f"$ {valor:,.0f}" if es_money else f"{valor:,.2f}"
    if isinstance(valor, int):
        return f"{valor:,}"
    return str(valor)


def pdf_reporte(
    titulo: str,
    subtitulo: str,
    header: list[str],
    rows: list[list],
    money_cols: Iterable[int] = (),
    horizontal: bool = True,
) -> bytes:
    buf = BytesIO()
    page = landscape(A4) if horizontal else A4
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        title=titulo,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloReporte", parent=estilos["Title"], fontSize=15, textColor=PRIMARIO
    )
    money_cols = set(money_cols)

    elementos = [Paragraph(titulo, estilo_titulo)]
    if subtitulo:
        elementos.append(Paragraph(subtitulo, estilos["Normal"]))
    elementos.append(Spacer(1, 8))

    datos = [header]
    for fila in rows:
        datos.append(
            [_pdf_celda(v, i in money_cols) for i, v in enumerate(fila)]
        )

    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PRIMARIO),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 9),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, CELDA_SUAVE],
                ),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elementos.append(tabla)
    doc.build(elementos)
    return buf.getvalue()


def excel_reporte(
    titulo: str,
    subtitulo: str,
    hojas: list[dict],
) -> bytes:
    """hojas: [{nombre, resumen: [(lbl, val)], header: [...], rows: [[...]], money_cols: {int}}]"""
    wb = Workbook()
    wb.remove(wb.active)
    for hoja in hojas:
        ws = wb.create_sheet(hoja["nombre"][:31])
        ws["A1"] = titulo
        ws["A1"].font = Font(bold=True, size=14, color=PRIMARIO_FILL)
        ws["A1"].fill = PatternFill("solid", fgColor=SUAVE_FILL)
        fila = 2
        if subtitulo:
            ws.cell(row=fila, column=1).value = subtitulo
            ws.cell(row=fila, column=1).font = Font(italic=True, color="666666")
            fila += 1
        for etiqueta, valor in hoja.get("resumen", []):
            ws.cell(row=fila, column=1).value = etiqueta
            ws.cell(row=fila, column=1).font = Font(bold=True)
            celda_valor = ws.cell(row=fila, column=2)
            celda_valor.value = valor
            celda_valor.number_format = "#,##0"
            fila += 1
        if hoja.get("resumen"):
            fila += 1
        header = hoja["header"]
        for col, nombre in enumerate(header):
            celda = ws.cell(row=fila, column=col + 1)
            celda.value = nombre
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=PRIMARIO_FILL)
        fila += 1
        money_cols = hoja.get("money_cols", ())
        for fila_datos in hoja["rows"]:
            for col, valor in enumerate(fila_datos):
                celda = ws.cell(row=fila, column=col + 1)
                if isinstance(valor, datetime):
                    valor = _sin_zona(valor)
                else:
                    valor = valor if valor is not None else ""
                celda.value = valor
                if col in money_cols and isinstance(valor, (int, float)):
                    celda.number_format = "#,##0"
            fila += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()