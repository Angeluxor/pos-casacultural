import openpyxl

wb = openpyxl.load_workbook('data/Sistema_Ventas_CasaCultural_MASTER.xlsx', read_only=True, data_only=True)

sheets_to_peek = ['Facturas', 'DetalleFactura', 'Pagos', 'Caja', 'MovimientosInventario']
for s in sheets_to_peek:
    ws = wb[s]
    print(f"\n=== {s} ===")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=4, values_only=True)):
        print(f"  Row {i}: {row}")

# Count rows
print("\n=== ROW COUNTS ===")
for s in ['Productos', 'Clientes', 'Facturas', 'DetalleFactura', 'Pagos', 'Caja', 'MovimientosInventario']:
    ws = wb[s]
    count = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None)
    print(f"  {s}: {count} rows")
